import datetime
import json
import logging
import io
import os
import pytz
from typing import List, Dict, Any, Tuple
from sqlalchemy.orm import Session
from sqlalchemy import desc

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.drawing.image import Image

from backend.models import (
    Student, Department, WeeklySession, 
    OfficialWeeklySnapshot, WeeklyStudentProgress, StudentStatSnapshot
)
from backend.services.weekly_session_resolver import resolve_weekly_sessions

logger = logging.getLogger(__name__)

def get_snapshot_data(db: Session, session_id: int) -> Tuple[Dict[str, dict], Dict[int, dict]]:
    """
    Returns (contest_snapshot_by_username, progress_by_student_id)
    """
    snapshot = db.query(OfficialWeeklySnapshot).filter(
        OfficialWeeklySnapshot.session_id == session_id
    ).first()
    
    contest_by_username = {}
    if snapshot and snapshot.dataset:
        data = snapshot.dataset
        if isinstance(data, str):
            data = json.loads(data)
        if isinstance(data, str):
            data = json.loads(data)
            
        rows = data.get("rows", [])
        for r in rows:
            uname = r.get("username")
            if uname:
                contest_by_username[uname] = r
                
    # Get week number from session
    session = db.query(WeeklySession).filter(WeeklySession.id == session_id).first()
    week_num = session.week_number if session else 0
    
    progress_by_student = {}
    if week_num:
        progs = db.query(WeeklyStudentProgress).filter(
            WeeklyStudentProgress.week_number == week_num
        ).all()
        for p in progs:
            progress_by_student[p.student_id] = p
            
    return contest_by_username, progress_by_student

def calculate_solved_bucket(total_solved: int) -> str:
    if total_solved > 500:
        return "Above 500"
    elif 250 <= total_solved <= 500:
        return "250–500"
    elif 100 <= total_solved < 250:
        return "Less than 250"
    elif 1 <= total_solved < 100:
        return "Less than 100"
    else:
        return "Not Yet Started"

def prefetch_global_ranks(db: Session, student_ids: List[int], date_limit: datetime.datetime) -> Dict[int, int]:
    """Finds the closest stat snapshot before or near the session end for all given students."""
    if not student_ids:
        return {}
        
    snaps = db.query(StudentStatSnapshot).filter(
        StudentStatSnapshot.student_id.in_(student_ids),
        StudentStatSnapshot.captured_at <= date_limit
    ).order_by(desc(StudentStatSnapshot.captured_at)).all()
    
    ranks = {}
    for snap in snaps:
        if snap.student_id not in ranks:
            ranks[snap.student_id] = snap.global_rank if snap.global_rank else 999999
    return ranks

def build_batch_report(db: Session, last_session_id: int, current_session_id: int, target_dept_id: int) -> Dict[str, Any]:
    last_contest, last_prog = get_snapshot_data(db, last_session_id)
    curr_contest, curr_prog = get_snapshot_data(db, current_session_id)
    
    last_session = db.query(WeeklySession).filter(WeeklySession.id == last_session_id).first()
    curr_session = db.query(WeeklySession).filter(WeeklySession.id == current_session_id).first()
    
    last_date = datetime.datetime.fromisoformat(last_session.session_date) if last_session else datetime.datetime.now(datetime.timezone.utc)
    curr_date = datetime.datetime.fromisoformat(curr_session.session_date) if curr_session else datetime.datetime.now(datetime.timezone.utc)
    
    students = db.query(Student).filter(
        Student.is_active == True,
        Student.department_id == target_dept_id
    ).all()
    
    student_ids = [s.id for s in students]
    last_ranks = prefetch_global_ranks(db, student_ids, last_date)
    curr_ranks = prefetch_global_ranks(db, student_ids, curr_date)
    
    batches = ["2023-2027", "2024-2028", "2025-2029"]
    report = {b: {"last": _empty_metrics(), "curr": _empty_metrics()} for b in batches}
    
    def process_week(st: Student, contest_map: dict, prog_map: dict, date_limit: datetime.datetime, is_last: bool):
        batch = st.batch
        if batch not in report:
            return
            
        m = report[batch]["last"] if is_last else report[batch]["curr"]
        m["total_students"] += 1
        
        # PRIMARY ONLY for Solved bucket
        prog = prog_map.get(st.id)
        primary_solved = prog.total_solved if prog else 0
        bucket = calculate_solved_bucket(primary_solved)
        m["solved"][bucket] += 1
        
        # CONTEST deduplication (PRIMARY + SECONDARY)
        p_uname = st.username
        s_uname = st.secondary_leetcode_id
        
        p_data = contest_map.get(p_uname) if p_uname else None
        s_data = contest_map.get(s_uname) if s_uname else None

        # Only count if status is PUBLIC
        p_public = p_data.get("status") == "PUBLIC" if p_data else False
        s_public = s_data.get("status") == "PUBLIC" if s_data else False
        
        q1 = bool(p_data.get("q1", 0) if p_public else 0) | bool(s_data.get("q1", 0) if s_public else 0)
        q2 = bool(p_data.get("q2", 0) if p_public else 0) | bool(s_data.get("q2", 0) if s_public else 0)
        q3 = bool(p_data.get("q3", 0) if p_public else 0) | bool(s_data.get("q3", 0) if s_public else 0)
        q4 = bool(p_data.get("q4", 0) if p_public else 0) | bool(s_data.get("q4", 0) if s_public else 0)
        
        total_q = q1 + q2 + q3 + q4
        if total_q == 4: m["contest"]["4 Q Solved"] += 1
        elif total_q == 3: m["contest"]["3 Q Solved"] += 1
        elif total_q == 2: m["contest"]["2 Q Solved"] += 1
        elif total_q == 1: m["contest"]["1 Q Solved"] += 1
        
        # RATING > 1500 (PRIMARY + SECONDARY)
        p_rating = prog.rating if prog and prog.rating else 0
        s_rating = s_data.get("contest_rating") or 0 if s_data else 0
        
        if (p_rating and p_rating > 1500) or (s_rating and s_rating > 1500):
            m["rating_above_1500"] += 1
            
        # RANKING < 20000 (PRIMARY + SECONDARY)
        p_rank = last_ranks.get(st.id, 999999) if is_last else curr_ranks.get(st.id, 999999)
        s_rank = s_data.get("contest_rank") or 999999 if s_data else 999999
        
        if (0 < p_rank < 20000) or (0 < s_rank < 20000):
            m["ranking_below_20000"] += 1

    for st in students:
        process_week(st, last_contest, last_prog, last_date, True)
        process_week(st, curr_contest, curr_prog, curr_date, False)
        
    return report

def _empty_metrics():
    return {
        "total_students": 0,
        "solved": {
            "Above 500": 0,
            "250–500": 0,
            "Less than 250": 0,
            "Less than 100": 0,
            "Not Yet Started": 0
        },
        "contest": {
            "4 Q Solved": 0,
            "3 Q Solved": 0,
            "2 Q Solved": 0,
            "1 Q Solved": 0
        },
        "rating_above_1500": 0,
        "ranking_below_20000": 0
    }

def get_target_sessions(db: Session):
    resolved = resolve_weekly_sessions(db)
    curr_sess = resolved.get("current_week_session")
    last_sess = resolved.get("last_week_session")
    
    if not curr_sess or not last_sess:
        tz = pytz.timezone('Asia/Kolkata')
        now = datetime.datetime.now(tz)
        
        sessions = db.query(WeeklySession).filter(
            WeeklySession.status == "FINALIZED",
            # Fallback check against IST boundary loosely if resolver fails
        ).order_by(WeeklySession.id.desc()).all()
        
        if len(sessions) < 2:
            raise ValueError("Need at least 2 finalized sessions to generate this report.")
            
        return sessions[0].id, sessions[1].id
        
    return curr_sess.id, last_sess.id

def _render_sheet(ws, dept_name: str, report_data: dict, current_session_id: int):
    # Print / Layout settings
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_area = "A1:M15"
    
    # 1. Title Rows
    ws.merge_cells("A1:M1")
    ws["A1"] = "NANDHA ENGINEERING COLLEGE, ERODE-52."
    ws["A1"].font = Font(bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center")
    
    ws.merge_cells("A2:M2")
    ws["A2"] = "DEPARTMENT OF " + dept_name.upper()
    ws["A2"].font = Font(bold=True, size=14)
    ws["A2"].alignment = Alignment(horizontal="center")
    
    ws.merge_cells("A3:M3")
    ws["A3"] = "LEETCODE WEEKLY PERFORMANCE REPORT"
    ws["A3"].font = Font(bold=True, size=14)
    ws["A3"].alignment = Alignment(horizontal="center")
    
    # Insert Logo if exists
    logo_path = os.path.join("backend", "static", "logo.png")
    if os.path.exists(logo_path):
        try:
            img = Image(logo_path)
            img.width = 100
            img.height = 100
            ws.add_image(img, "A1")
        except Exception as e:
            logger.warning(f"Failed to add logo: {e}")
            
    # 2. Table Headers (A to M -> exactly 13 columns)
    headers = [
        "Batch", "Strength", 
        "Above 500", "250–500", "Less than 250", "Less than 100", "Not Yet Started",
        "4 Q Solved", "3 Q Solved", "2 Q Solved", "1 Q Solved",
        "Rating (>1500)", "Ranking (<20K)"
    ]
    
    # Merge for main categories
    ws.merge_cells("C5:G5")
    ws["C5"] = "Problems Solved Categories"
    ws["C5"].alignment = Alignment(horizontal="center")
    ws["C5"].font = Font(bold=True)
    ws["C5"].fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    
    ws.merge_cells("H5:K5")
    ws["H5"] = "Contest Questions Solved"
    ws["H5"].alignment = Alignment(horizontal="center")
    ws["H5"].font = Font(bold=True)
    ws["H5"].fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    
    header_row = 6
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col, value=h)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        
    # 3. Fill Data
    row_num = 7
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                         top=Side(style='thin'), bottom=Side(style='thin'))
                         
    for batch_idx, batch in enumerate(["2023-2027", "2024-2028", "2025-2029"]):
        batch_data = report_data.get(batch, {})
        for week_label, week_name in [("last", "Last Week"), ("curr", "Current Week")]:
            m = batch_data.get(week_label, _empty_metrics())
            
            # Batch
            ws.cell(row=row_num, column=1, value=f"{batch} ({week_name})")
            ws.cell(row=row_num, column=2, value=m["total_students"])
            
            # Solved
            ws.cell(row=row_num, column=3, value=m["solved"]["Above 500"])
            ws.cell(row=row_num, column=4, value=m["solved"]["250–500"])
            ws.cell(row=row_num, column=5, value=m["solved"]["Less than 250"])
            ws.cell(row=row_num, column=6, value=m["solved"]["Less than 100"])
            ws.cell(row=row_num, column=7, value=m["solved"]["Not Yet Started"])
            
            # Contest
            ws.cell(row=row_num, column=8, value=m["contest"]["4 Q Solved"])
            ws.cell(row=row_num, column=9, value=m["contest"]["3 Q Solved"])
            ws.cell(row=row_num, column=10, value=m["contest"]["2 Q Solved"])
            ws.cell(row=row_num, column=11, value=m["contest"]["1 Q Solved"])
            
            # Rating & Ranking
            ws.cell(row=row_num, column=12, value=m["rating_above_1500"])
            ws.cell(row=row_num, column=13, value=m["ranking_below_20000"])
            
            for col in range(1, 14):
                ws.cell(row=row_num, column=col).border = thin_border
                ws.cell(row=row_num, column=col).alignment = Alignment(horizontal="center")
                
            row_num += 1
            
    # Adjust column widths
    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 12
    for col in ["C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M"]:
        ws.column_dimensions[col].width = 12

def generate_weekly_performance_excel(db: Session, target_dept_name: str = None) -> bytes:
    current_session_id, last_session_id = get_target_sessions(db)
    
    wb = Workbook()
    # Remove default sheet
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])
        
    departments = ["Cyber Security", "Internet of Things"]
    
    for dept_name in departments:
        dept = db.query(Department).filter(Department.name == dept_name).first()
        if not dept:
            continue
            
        report_data = build_batch_report(db, last_session_id, current_session_id, dept.id)
        
        ws = wb.create_sheet(title=dept_name)
        _render_sheet(ws, dept_name, report_data, current_session_id)
        
    if not wb.sheetnames:
        ws = wb.create_sheet("Empty")
        ws["A1"] = "No matching departments found."
        
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
