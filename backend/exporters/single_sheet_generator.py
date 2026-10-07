import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from typing import List, Dict, Any, Optional
import datetime
from backend.exporters.nec_master_excel_design import apply_master_college_identity
from backend.models import Student, WeeklyPublicResult, WeeklyVirtualResult
from backend.services.weekly_session_resolver import resolve_target_weekly_session

def generate_single_sheet_contest_excel(
    db: Session,
    dataset: Dict[str, Any],
    session_id: str
) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    if not ws:
        ws = wb.create_sheet()
    ws.title = "Contest Report"
    
    target_session = resolve_target_weekly_session(db, session_id)
    public_map = {}
    virtual_map = {}
    if target_session:
        p_list = db.query(WeeklyPublicResult, Student).join(Student, WeeklyPublicResult.student_id == Student.id).filter(WeeklyPublicResult.session_id == target_session.id).all()
        for pr, st in p_list:
            public_map[str(st.reg_no)] = pr
        v_list = db.query(WeeklyVirtualResult, Student).join(Student, WeeklyVirtualResult.student_id == Student.id).filter(WeeklyVirtualResult.session_id == target_session.id).all()
        for vr, st in v_list:
            virtual_map[str(st.reg_no)] = vr

    # Extract data
    students = (
        dataset.get("allStudents") or 
        dataset.get("rows") or 
        dataset.get("all_students_current") or 
        dataset.get("all_students") or 
        dataset.get("all_rows") or 
        dataset.get("students") or 
        []
    )
    
    contest_name = str(dataset.get("contest_name") or dataset.get("contestName") or dataset.get("title") or "")
    contest_date = str(dataset.get("session_date") or dataset.get("sessionDate") or dataset.get("report_date") or "")
    dept = dataset.get("deptFilter") or dataset.get("department") or "ALL"
    year = dataset.get("yearFilter") or dataset.get("year") or "ALL"
    
    # Apply Master Identity Header
    next_row = apply_master_college_identity(
        ws=ws,
        report_title="Student Contest Performance Single Sheet",
        department=dept,
        year=year,
        contest_name=contest_name,
        session_date=contest_date,
        total_roster=len(students),
        cols=15
    )
    
    # Fonts and Alignments
    font_header = Font(name="Times New Roman", size=10, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    font_data = Font(name="Times New Roman", size=10)
    font_name = Font(name="Times New Roman", size=10, bold=True)
    
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=False)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=False)
    
    thin_side = Side(style='thin', color='000000')
    grid_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    
    headers = [
        "S.No", 
        "Reg No", 
        "Name", 
        "Leetcode ID", 
        "Q1", 
        "Q2", 
        "Q3", 
        "Q4", 
        "Contest Name", 
        "Contest Date", 
        "Total Solved", 
        "Contest Easy", 
        "Medium", 
        "Hard", 
        "Fetch Details/Error (Time)"
    ]
    
    # Write Table Header
    ws.row_dimensions[next_row].height = 25
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=next_row, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = grid_border
        
    row_idx = next_row + 1
    
    for s_idx, s in enumerate(students, 1):
        reg_no = str(s.get("reg_no") or s.get("register_no") or "")
        pr = public_map.get(reg_no)
        vr = virtual_map.get(reg_no)

        q1 = q2 = q3 = q4 = 0
        if pr:
            q1 = 1 if pr.q1 else 0
            q2 = 1 if pr.q2 else 0
            q3 = 1 if pr.q3 else 0
            q4 = 1 if pr.q4 else 0
        elif vr:
            q1 = 1 if vr.q1 else 0
            q2 = 1 if vr.q2 else 0
            q3 = 1 if vr.q3 else 0
            q4 = 1 if vr.q4 else 0
        else:
            q1 = 1 if str(s.get("q1", "")).lower() in ("1", "true") else 0
            q2 = 1 if str(s.get("q2", "")).lower() in ("1", "true") else 0
            q3 = 1 if str(s.get("q3", "")).lower() in ("1", "true") else 0
            q4 = 1 if str(s.get("q4", "")).lower() in ("1", "true") else 0
        
        # Calculate difficulties strictly based on questions actually solved
        c_easy = q1
        c_med = q2 + q3
        c_hard = q4
        
        # Determine fetch time/error string
        fetch_time = ""
        st_val = str(s.get("status") or "").upper()
        if st_val in ["DATA_ERROR", "TIMEOUT", "SOURCE_UNAVAILABLE", "FETCH_FAILED", "USERNAME_NOT_FOUND", "ERROR"]:
            fetch_time = f"ERROR - {st_val}"
        else:
            ist_tz = datetime.timezone(datetime.timedelta(hours=5, minutes=30))
            if s.get("fetched_at"):
                try:
                    val = s.get("fetched_at")
                    if isinstance(val, str):
                        dt = datetime.datetime.fromisoformat(val.replace("Z", "+00:00"))
                        fetch_time = "SUCCESS - " + dt.astimezone(ist_tz).strftime("%I:%M %p")
                    else:
                        fetch_time = "SUCCESS - " + val.astimezone(ist_tz).strftime("%I:%M %p")
                except:
                    fetch_time = "SUCCESS - " + datetime.datetime.now(ist_tz).strftime("%I:%M %p") 
            else:
                fetch_time = "SUCCESS - " + datetime.datetime.now(ist_tz).strftime("%I:%M %p") 
        
        tot = s.get("solved")
        if not tot:
            tot = s.get("total_solved")
        if not tot:
            tot = s.get("total_contest_solved")
        
        try:
            tot = int(tot) if tot is not None else 0
        except:
            tot = 0
            
        vals = [
            s_idx,
            reg_no,
            str(s.get("name") or s.get("student_name") or ""),
            str(s.get("username") or s.get("leetcode_username") or ""),
            q1,
            q2,
            q3,
            q4,
            contest_name,
            contest_date,
            max(tot, q1+q2+q3+q4),
            c_easy,
            c_med,
            c_hard,
            fetch_time
        ]
        
        for col_idx, val in enumerate(vals, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_name if col_idx == 3 else font_data
            cell.alignment = align_left if col_idx == 3 else align_center
            cell.border = grid_border
            
        ws.row_dimensions[row_idx].height = 20
        row_idx += 1
        
    # Apply column widths
    widths = [8, 15, 30, 20, 8, 8, 8, 8, 20, 15, 12, 12, 12, 12, 30]
    for col_idx, width in enumerate(widths, 1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width
        
    from io import BytesIO
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
