"""
original_contest_excel_exporter.py
================================================================================
LEETCODE INSTITUTIONAL ORIGINAL CONTEST EXCEL REPORT GENERATOR
================================================================================
Generates the authoritative 6-sheet Excel workbook for Original Contest Result
Reconstruction adhering strictly to Times New Roman typography, Navy headers,
institutional formatting, precise column alignment, clean borders, and proper spacing.
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import Dict, Any, List

# Master Institutional Style Constants (Times New Roman)
FONT_TNR = "Times New Roman"

NAVY_HEADER_FILL = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
NAVY_SUBHEADER_FILL = PatternFill(start_color="2E5B88", end_color="2E5B88", fill_type="solid")
ALT_ROW_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
SECTION_FILL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

FILL_GREEN = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
FONT_GREEN = Font(name=FONT_TNR, size=10, bold=True, color="047857")

FILL_AMBER = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
FONT_AMBER = Font(name=FONT_TNR, size=10, bold=True, color="B45309")

FILL_RED = PatternFill(start_color="FFF1F2", end_color="FFF1F2", fill_type="solid")
FONT_RED = Font(name=FONT_TNR, size=10, bold=True, color="B91C1C")

FONT_HEADER = Font(name=FONT_TNR, size=10, bold=True, color="FFFFFF")
FONT_BODY = Font(name=FONT_TNR, size=10)
FONT_BODY_BOLD = Font(name=FONT_TNR, size=10, bold=True)
FONT_TITLE = Font(name=FONT_TNR, size=13, bold=True, color="FFFFFF")
FONT_SUBTITLE = Font(name=FONT_TNR, size=9, italic=True, color="FFFFFF")

ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

_THIN_SIDE = Side(style="thin", color="CBD5E1")
_THIN_BORDER = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE)


def apply_auto_column_widths_and_borders(ws):
    """
    Applies explicit borders to every cell and dynamically calculates optimal column widths.
    """
    ws.sheet_view.showGridLines = True
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        header_val = str(ws.cell(row=1, column=col[0].column).value or ws.cell(row=4, column=col[0].column).value or "")

        for cell in col:
            # Ensure every single cell gets thin border
            cell.border = _THIN_BORDER
            val_str = str(cell.value or "")
            if val_str.startswith("http"):
                val_str = "https://leetcode.com/u/user/"
            max_len = max(max_len, len(val_str))

        # Precision width assignment
        if "s.no" in header_val.lower():
            ws.column_dimensions[col_letter].width = 8
        elif "registration number" in header_val.lower():
            ws.column_dimensions[col_letter].width = 20
        elif header_val in ("Year", "Q1", "Q2", "Q3", "Q4"):
            ws.column_dimensions[col_letter].width = 9
        elif "department" in header_val.lower():
            ws.column_dimensions[col_letter].width = 14
        elif "verified" in header_val.lower():
            ws.column_dimensions[col_letter].width = 14
        else:
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 45)


def generate_original_contest_excel_report(
    reconstruction_payload: Dict[str, Any],
    output_filepath: str
) -> str:
    """
    Generates the complete 6-sheet Excel workbook from reconstruction payload with clean borders and alignment.
    """
    wb = openpyxl.Workbook()
    if wb.active is not None:
        wb.remove(wb.active)  # Remove default sheet

    meta = reconstruction_payload.get("snapshot_metadata", {})

    # -------------------------------------------------------------------------
    # SHEET 1: OFFICIAL_RESULT
    # -------------------------------------------------------------------------
    ws1 = wb.create_sheet(title="OFFICIAL_RESULT")
    
    headers1 = [
        "S.No", "Registration Number", "Student Name", "Department", "Year",
        "LeetCode Username", "LeetCode Profile", "Account Verification", "Participation Type",
        "Contest ID", "Contest Name", "Contest Date",
        "Q1 Problem", "Q1 Status", "Q1 Evidence", "Q1 Submission Time",
        "Q2 Problem", "Q2 Status", "Q2 Evidence", "Q2 Submission Time",
        "Q3 Problem", "Q3 Status", "Q3 Evidence", "Q3 Submission Time",
        "Q4 Problem", "Q4 Status", "Q4 Evidence", "Q4 Submission Time",
        "Verified Q1", "Verified Q2", "Verified Q3", "Verified Q4",
        "Verified Contest Solved", "Verification Status"
    ]
    num_cols1 = len(headers1)
    end_col_letter1 = get_column_letter(num_cols1)

    # Metadata Title Banner
    ws1.merge_cells(f"A1:{end_col_letter1}1")
    title_cell = ws1["A1"]
    title_cell.value = f"NANDHA ENGINEERING COLLEGE — {meta.get('report_type', 'ORIGINAL VERIFIED CONTEST RESULT')}"
    title_cell.font = FONT_TITLE
    title_cell.fill = NAVY_HEADER_FILL
    title_cell.alignment = ALIGN_CENTER
    ws1.row_dimensions[1].height = 28

    ws1.merge_cells(f"A2:{end_col_letter1}2")
    sub_cell = ws1["A2"]
    sub_cell.value = f"Contest: {meta.get('contest_name')} ({meta.get('contest_id')}) | Date: {meta.get('contest_date')} | Window: {meta.get('contest_window')} | Snapshot: {meta.get('snapshot_id')} | Generated: {meta.get('generated_at')}"
    sub_cell.font = FONT_SUBTITLE
    sub_cell.fill = NAVY_HEADER_FILL
    sub_cell.alignment = ALIGN_CENTER
    ws1.row_dimensions[2].height = 18

    # Blank Row 3 with borders
    ws1.append([""] * num_cols1)
    ws1.row_dimensions[3].height = 10

    # Header Row 4
    ws1.append(headers1)
    ws1.row_dimensions[4].height = 26

    for col_num in range(1, num_cols1 + 1):
        cell = ws1.cell(row=4, column=col_num)
        cell.font = FONT_HEADER
        cell.fill = NAVY_SUBHEADER_FILL
        cell.alignment = ALIGN_CENTER
        cell.border = _THIN_BORDER

    # Center-aligned column indices for Sheet 1
    center_cols_s1 = {1, 2, 4, 5, 6, 8, 9, 10, 12, 14, 16, 18, 20, 22, 24, 26, 28, 29, 30, 31, 32, 33, 34}

    for row_idx, r in enumerate(reconstruction_payload.get("official_result", []), start=5):
        row_data = [
            r["s_no"], r["reg_no"], r["student_name"], r["department"], r["year"],
            r["leetcode_username"], r["leetcode_profile"], r["account_verification"], r["participation_type"],
            r["contest_id"], r["contest_name"], r["contest_date"],
            r["q1_problem"], r["q1_status"], r["q1_evidence"], r["q1_sub_time"],
            r["q2_problem"], r["q2_status"], r["q2_evidence"], r["q2_sub_time"],
            r["q3_problem"], r["q3_status"], r["q3_evidence"], r["q3_sub_time"],
            r["q4_problem"], r["q4_status"], r["q4_evidence"], r["q4_sub_time"],
            r["verified_q1"], r["verified_q2"], r["verified_q3"], r["verified_q4"],
            r["verified_total"], r["verification_status"]
        ]
        ws1.append(row_data)
        ws1.row_dimensions[row_idx].height = 22

        is_alt = (row_idx % 2 == 0)
        for col_num in range(1, len(row_data) + 1):
            cell = ws1.cell(row=row_idx, column=col_num)
            cell.font = FONT_BODY
            cell.border = _THIN_BORDER
            
            # Alignments
            if col_num in center_cols_s1:
                cell.alignment = ALIGN_CENTER
            else:
                cell.alignment = ALIGN_LEFT

            if is_alt:
                cell.fill = ALT_ROW_FILL

            # Status highlights
            if col_num in (29, 30, 31, 32, 33):  # Verified flags & total
                cell.font = FONT_BODY_BOLD
                if r["verified_total"] > 0:
                    cell.fill = FILL_GREEN
            elif col_num == 34:  # Verification Status
                if r["verification_status"].startswith("VERIFIED"):
                    cell.fill = FILL_GREEN
                    cell.font = FONT_GREEN
                elif r["verification_status"] in ("MISSING_USERNAME", "DUPLICATE_ACCOUNT_FLAGGED"):
                    cell.fill = FILL_RED
                    cell.font = FONT_RED

    apply_auto_column_widths_and_borders(ws1)

    # -------------------------------------------------------------------------
    # SHEET 2: EVIDENCE_AUDIT
    # -------------------------------------------------------------------------
    ws2 = wb.create_sheet(title="EVIDENCE_AUDIT")
    headers2 = [
        "Registration Number", "Student Name", "LeetCode Username", "Contest ID",
        "Problem ID", "Question Number", "Submission ID", "Submission Status",
        "Submission Timestamp UTC", "Submission Timestamp IST", "Participation Type",
        "Inside Contest Window", "Current Contest Match", "Current Problem Match",
        "Actual Participation", "Verified", "Counted", "Rejection Reason",
        "Evidence Source", "Snapshot ID"
    ]
    ws2.append(headers2)
    ws2.row_dimensions[1].height = 26

    for col_num in range(1, len(headers2) + 1):
        cell = ws2.cell(row=1, column=col_num)
        cell.font = FONT_HEADER
        cell.fill = NAVY_HEADER_FILL
        cell.alignment = ALIGN_CENTER
        cell.border = _THIN_BORDER

    center_cols_s2 = {1, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17}

    for row_idx, r in enumerate(reconstruction_payload.get("evidence_audit", []), start=2):
        row_data = [
            r["reg_no"], r["student_name"], r["leetcode_username"], r["contest_id"],
            r["problem_id"], r["question_number"], r["submission_id"], r["submission_status"],
            r["submission_timestamp_utc"], r["submission_timestamp_ist"], r["participation_type"],
            r["is_within_contest_window"], r["is_contest_problem"], r["is_contest_problem"],
            r["is_actual_participation"], r["is_verified"], r["counted"], r["rejection_reason"],
            r["evidence_source"], r["snapshot_id"]
        ]
        ws2.append(row_data)
        ws2.row_dimensions[row_idx].height = 20

        for col_num in range(1, len(row_data) + 1):
            cell = ws2.cell(row=row_idx, column=col_num)
            cell.font = FONT_BODY
            cell.border = _THIN_BORDER
            
            if col_num in center_cols_s2:
                cell.alignment = ALIGN_CENTER
            else:
                cell.alignment = ALIGN_LEFT

            if r["counted"]:
                if col_num in (16, 17):
                    cell.fill = FILL_GREEN
                    cell.font = FONT_GREEN

    apply_auto_column_widths_and_borders(ws2)

    # -------------------------------------------------------------------------
    # SHEET 3: STUDENT_SUMMARY
    # -------------------------------------------------------------------------
    ws3 = wb.create_sheet(title="STUDENT_SUMMARY")
    headers3 = [
        "S.No", "Registration Number", "Student Name", "Department", "Year",
        "LeetCode Username", "Participation Type", "Q1", "Q2", "Q3", "Q4",
        "Verified Total", "Verification Status"
    ]
    ws3.append(headers3)
    ws3.row_dimensions[1].height = 26

    for col_num in range(1, len(headers3) + 1):
        cell = ws3.cell(row=1, column=col_num)
        cell.font = FONT_HEADER
        cell.fill = NAVY_HEADER_FILL
        cell.alignment = ALIGN_CENTER
        cell.border = _THIN_BORDER

    center_cols_s3 = {1, 2, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13}

    for row_idx, r in enumerate(reconstruction_payload.get("student_summary", []), start=2):
        row_data = [
            row_idx - 1, r["reg_no"], r["student_name"], r["department"], r["year"],
            r["leetcode_username"], r["participation_type"],
            r["q1"], r["q2"], r["q3"], r["q4"],
            r["verified_total"], r["verification_status"]
        ]
        ws3.append(row_data)
        ws3.row_dimensions[row_idx].height = 22

        is_alt = (row_idx % 2 == 0)
        for col_num in range(1, len(row_data) + 1):
            cell = ws3.cell(row=row_idx, column=col_num)
            cell.font = FONT_BODY
            cell.border = _THIN_BORDER

            if col_num in center_cols_s3:
                cell.alignment = ALIGN_CENTER
            else:
                cell.alignment = ALIGN_LEFT

            if is_alt:
                cell.fill = ALT_ROW_FILL

            if col_num in (8, 9, 10, 11, 12):
                cell.font = FONT_BODY_BOLD
                if r["verified_total"] > 0:
                    cell.fill = FILL_GREEN

    apply_auto_column_widths_and_borders(ws3)

    # -------------------------------------------------------------------------
    # SHEET 4: RECONCILIATION
    # -------------------------------------------------------------------------
    ws4 = wb.create_sheet(title="RECONCILIATION")
    headers4 = [
        "S.No", "Registration Number", "Student Name", "LeetCode Username",
        "Old Q1", "Old Q2", "Old Q3", "Old Q4", "Old Total",
        "New Q1", "New Q2", "New Q3", "New Q4", "New Verified Total",
        "Difference", "Reason for Difference", "Evidence Status"
    ]
    ws4.append(headers4)
    ws4.row_dimensions[1].height = 26

    for col_num in range(1, len(headers4) + 1):
        cell = ws4.cell(row=1, column=col_num)
        cell.font = FONT_HEADER
        cell.fill = NAVY_HEADER_FILL
        cell.alignment = ALIGN_CENTER
        cell.border = _THIN_BORDER

    center_cols_s4 = {1, 2, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 17}

    for row_idx, r in enumerate(reconstruction_payload.get("reconciliation", []), start=2):
        row_data = [
            row_idx - 1, r["reg_no"], r["student_name"], r["username"],
            r["old_q1"], r["old_q2"], r["old_q3"], r["old_q4"], r["old_total"],
            r["new_q1"], r["new_q2"], r["new_q3"], r["new_q4"], r["new_verified_total"],
            r["difference"], r["reason_for_difference"], r["evidence_status"]
        ]
        ws4.append(row_data)
        ws4.row_dimensions[row_idx].height = 22

        for col_num in range(1, len(row_data) + 1):
            cell = ws4.cell(row=row_idx, column=col_num)
            cell.font = FONT_BODY
            cell.border = _THIN_BORDER

            if col_num in center_cols_s4:
                cell.alignment = ALIGN_CENTER
            else:
                cell.alignment = ALIGN_LEFT

            if col_num == 15:  # Difference
                cell.font = FONT_BODY_BOLD
                if r["difference"] != 0:
                    cell.fill = FILL_AMBER
                    cell.font = FONT_AMBER
                else:
                    cell.fill = FILL_GREEN

    apply_auto_column_widths_and_borders(ws4)

    # -------------------------------------------------------------------------
    # SHEET 5: STATISTICS
    # -------------------------------------------------------------------------
    ws5 = wb.create_sheet(title="STATISTICS")
    ws5.sheet_view.showGridLines = True
    stats = reconstruction_payload.get("statistics", {})

    ws5.append(["NANDHA ENGINEERING COLLEGE — CONTEST RECONSTRUCTION STATISTICS"])
    ws5["A1"].font = FONT_TITLE
    ws5["A1"].fill = NAVY_HEADER_FILL
    ws5.merge_cells("A1:D1")
    ws5.row_dimensions[1].height = 28

    ws5.append([""] * 4)

    # Table 1: Execution Overview
    ws5.append(["Metric", "Count", "Percentage", "Status"])
    ws5.row_dimensions[3].height = 24
    for c in range(1, 5):
        cell = ws5.cell(row=3, column=c)
        cell.font = FONT_HEADER
        cell.fill = NAVY_SUBHEADER_FILL
        cell.alignment = ALIGN_CENTER
        cell.border = _THIN_BORDER

    tot_s = max(1, stats.get("total_students", 1))
    t1_rows = [
        ("Total Roster Students Evaluated", stats.get("total_students", 0), "100.0%", "VERIFIED"),
        ("ACTUAL Contest Participants", stats.get("actual_participants", 0), f"{round(stats.get('actual_participants', 0) / tot_s * 100, 1)}%", "VERIFIED"),
        ("VIRTUAL Practice Participants", stats.get("virtual_participants", 0), f"{round(stats.get('virtual_participants', 0) / tot_s * 100, 1)}%", "EXCLUDED"),
        ("NOT_VERIFIED / Absent", stats.get("not_verified_participants", 0), f"{round(stats.get('not_verified_participants', 0) / tot_s * 100, 1)}%", "AUDITED"),
    ]
    for r in t1_rows:
        ws5.append(list(r))

    ws5.append([""] * 4)

    # Table 2: Solve Distribution (0, 1, 2, 3, 4 Solved)
    ws5.append(["Performance Tier", "Verified Solved Count", "Student Count", "Percentage"])
    ws5.row_dimensions[9].height = 24
    for c in range(1, 5):
        cell = ws5.cell(row=9, column=c)
        cell.font = FONT_HEADER
        cell.fill = NAVY_SUBHEADER_FILL
        cell.alignment = ALIGN_CENTER
        cell.border = _THIN_BORDER

    t2_rows = [
        ("4/4 Perfect Solved", "4 / 4", stats.get("s4_count", 0), f"{round(stats.get('s4_count', 0) / tot_s * 100, 1)}%"),
        ("3/4 Solved", "3 / 4", stats.get("s3_count", 0), f"{round(stats.get('s3_count', 0) / tot_s * 100, 1)}%"),
        ("2/4 Solved", "2 / 4", stats.get("s2_count", 0), f"{round(stats.get('s2_count', 0) / tot_s * 100, 1)}%"),
        ("1/4 Solved", "1 / 4", stats.get("s1_count", 0), f"{round(stats.get('s1_count', 0) / tot_s * 100, 1)}%"),
        ("0/4 Solved", "0 / 4", stats.get("s0_count", 0), f"{round(stats.get('s0_count', 0) / tot_s * 100, 1)}%"),
    ]
    for r in t2_rows:
        ws5.append(list(r))

    ws5.append([""] * 4)

    # Table 3: Question-Level Solve Totals
    ws5.append(["Question", "Verified Solved Students", "", ""])
    ws5.row_dimensions[16].height = 24
    for c in range(1, 3):
        cell = ws5.cell(row=16, column=c)
        cell.font = FONT_HEADER
        cell.fill = NAVY_SUBHEADER_FILL
        cell.alignment = ALIGN_CENTER
        cell.border = _THIN_BORDER

    ws5.append(["Q1 Solved Total", stats.get("q1_solved_count", 0), "", ""])
    ws5.append(["Q2 Solved Total", stats.get("q2_solved_count", 0), "", ""])
    ws5.append(["Q3 Solved Total", stats.get("q3_solved_count", 0), "", ""])
    ws5.append(["Q4 Solved Total", stats.get("q4_solved_count", 0), "", ""])
    ws5.append(["Total Verified Solves (Sum Q1..Q4)", stats.get("total_verified_solves", 0), "", ""])

    apply_auto_column_widths_and_borders(ws5)

    # -------------------------------------------------------------------------
    # SHEET 6: DATA_QUALITY
    # -------------------------------------------------------------------------
    ws6 = wb.create_sheet(title="DATA_QUALITY")
    headers6 = [
        "S.No", "Registration Number", "Student Name", "LeetCode Username",
        "Issue Category", "Issue Description", "Impact on Score", "Action Required"
    ]
    ws6.append(headers6)
    ws6.row_dimensions[1].height = 26

    for col_num in range(1, len(headers6) + 1):
        cell = ws6.cell(row=1, column=col_num)
        cell.font = FONT_HEADER
        cell.fill = NAVY_HEADER_FILL
        cell.alignment = ALIGN_CENTER
        cell.border = _THIN_BORDER

    center_cols_s6 = {1, 2, 4, 5}

    for row_idx, r in enumerate(reconstruction_payload.get("data_quality", []), start=2):
        row_data = [
            row_idx - 1, r["reg_no"], r["student_name"], r["leetcode_username"],
            r["issue_category"], r["issue_description"], r["impact_on_score"], r["action_required"]
        ]
        ws6.append(row_data)
        ws6.row_dimensions[row_idx].height = 22

        for col_num in range(1, len(row_data) + 1):
            cell = ws6.cell(row=row_idx, column=col_num)
            cell.font = FONT_BODY
            cell.border = _THIN_BORDER

            if col_num in center_cols_s6:
                cell.alignment = ALIGN_CENTER
            else:
                cell.alignment = ALIGN_LEFT

            if col_num == 5:
                cell.font = FONT_RED
                cell.fill = FILL_RED

    apply_auto_column_widths_and_borders(ws6)

    # Ensure reports directory exists
    os.makedirs(os.path.dirname(output_filepath), exist_ok=True)
    wb.save(output_filepath)
    return output_filepath
