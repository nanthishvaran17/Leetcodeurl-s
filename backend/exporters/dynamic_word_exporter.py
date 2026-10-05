import io
import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT

def export_dynamic_word(dataset: dict) -> bytes:
    doc = Document()
    
    # Set to Landscape and narrow margins to prevent word-wrapping in large tables
    sections = doc.sections
    for section in sections:
        new_width, new_height = section.page_height, section.page_width
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = new_width
        section.page_height = new_height
        section.left_margin = Inches(0.3)
        section.right_margin = Inches(0.3)
        section.top_margin = Inches(0.4)
        section.bottom_margin = Inches(0.4)
    
    # Metadata & WoW detection
    report_type = str(dataset.get("reportType") or dataset.get("report_type") or "").upper().strip()
    is_wow = (
        report_type in ("WEEK_ON_WEEK_INTELLIGENCE", "WOW_INTEL", "WEEK_ON_WEEK")
        or (rows and any(k in rows[0] for k in ("prev_status", "curr_status", "solved_delta", "trend")))
        or "wowSummary" in dataset
    )

    prev_c_lbl = str(dataset.get("prevContest") or (dataset.get("wowSummary") or {}).get("prevContest") or "Last Week").strip()
    curr_c_lbl = str(dataset.get("currContest") or (dataset.get("wowSummary") or {}).get("currContest") or dataset.get("contestName") or "This Week").strip()

    if is_wow:
        report_title = f"WEEK-ON-WEEK INTELLIGENCE — {prev_c_lbl.upper()} VS {curr_c_lbl.upper()}"
    else:
        report_title = dataset.get("reportTitle") or "Dynamic Report"

    title = doc.add_heading(report_title.upper(), level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = RGBColor(27, 54, 93)
        run.font.size = Pt(14)
        
    doc.add_paragraph() # Spacer
    
    rows = dataset.get("rows") or dataset.get("allStudents") or dataset.get("all_students_current") or []
    if not rows:
        doc.add_paragraph("No data available for this report.")
        output = io.BytesIO()
        doc.save(output)
        return output.getvalue()
        
    # Extract headers
    first_row = rows[0]
    
    if is_wow:
        clean_headers = [
            "S.No",
            "Register No",
            "Student Name",
            "Dept",
            "Yr",
            f"{prev_c_lbl} Status",
            f"{prev_c_lbl} Solved",
            f"{prev_c_lbl} Score",
            f"{curr_c_lbl} Status",
            f"{curr_c_lbl} Solved",
            f"{curr_c_lbl} Score",
            "Δ Solved",
            "Trend",
        ]
        num_cols = len(clean_headers)
        table = doc.add_table(rows=1, cols=num_cols)
        table.style = 'Table Grid'
        table.autofit = True
        
        # Write headers
        hdr_cells = table.rows[0].cells
        for i, h in enumerate(clean_headers):
            hdr_cells[i].text = h
            for run in hdr_cells[i].paragraphs[0].runs:
                run.font.bold = True
                run.font.size = Pt(7.5)
                
        # Write data rows
        for idx, r in enumerate(rows, 1):
            row_cells = table.add_row().cells
            sno_val = str(r.get("s_no") or idx)
            reg_val = str(r.get("reg_no") or "")
            name_val = str(r.get("name") or "")
            dept_val = str(r.get("dept") or "")
            yr_val = str(r.get("year") or "")

            p_st_raw = str(r.get("prev_status") or "").upper().strip()
            p_st = "ATTENDED" if p_st_raw in ("ATTENDED", "PUBLIC", "VIRTUAL") else "ABSENT"
            p_sol = "—" if p_st == "ABSENT" and (r.get("prev_solved") is None or r.get("prev_solved") == 0) else str(r.get("prev_solved", "0"))
            p_sc = "—" if p_st == "ABSENT" and (r.get("prev_score") is None or r.get("prev_score") == 0) else str(r.get("prev_score", "0"))

            c_st_raw = str(r.get("curr_status") or "").upper().strip()
            c_st = "ATTENDED" if c_st_raw in ("ATTENDED", "PUBLIC", "VIRTUAL") else "ABSENT"
            c_sol = str(r.get("curr_solved", "0"))
            c_sc = str(r.get("curr_score", "0"))

            diff = r.get("solved_delta", 0)
            diff_str = f"+{diff}" if isinstance(diff, int) and diff > 0 else str(diff)
            trend_str = str(r.get("trend") or "Stable →")

            row_data = [sno_val, reg_val, name_val, dept_val, yr_val, p_st, p_sol, p_sc, c_st, c_sol, c_sc, diff_str, trend_str]
            for c_idx, val in enumerate(row_data):
                row_cells[c_idx].text = str(val)
                for run in row_cells[c_idx].paragraphs[0].runs:
                    run.font.size = Pt(7.5)
    else:
        weekly_data_keys = []
        if "weeklyData" in first_row:
            for wd in first_row.get("weeklyData", []):
                weekly_data_keys.append(wd.get("contestNum"))

        clean_headers = []
        headers_keys = list(first_row.keys())
        
        for h in headers_keys:
            if h == "weeklyData":
                for c_num in weekly_data_keys:
                    clean_headers.append(f"C{c_num}")
            else:
                formatted_h = re.sub(r'([a-z])([A-Z])', r'\1 \2', h).replace("_", " ").title().replace('Pct', '%')
                clean_headers.append(formatted_h)
                
        num_cols = len(clean_headers)
        table = doc.add_table(rows=1, cols=num_cols)
        table.style = 'Table Grid'
        table.autofit = True
        
        # Write headers
        hdr_cells = table.rows[0].cells
        for i, h in enumerate(clean_headers):
            hdr_cells[i].text = h
            for run in hdr_cells[i].paragraphs[0].runs:
                run.font.bold = True
                run.font.size = Pt(7) # Tiny font so it doesn't wrap
                
        # Write data rows
        for r in rows:
            row_cells = table.add_row().cells
            c_idx = 0
            for h in headers_keys:
                if h == "weeklyData":
                    wd_list = r.get("weeklyData") or []
                    wd_map = {wd.get("contestNum"): wd for wd in wd_list}
                    for c_num in weekly_data_keys:
                        wd_val = wd_map.get(c_num)
                        val = wd_val.get("solved", 0) if wd_val and wd_val.get("att") else "-"
                        row_cells[c_idx].text = str(val)
                        for run in row_cells[c_idx].paragraphs[0].runs:
                            run.font.size = Pt(7)
                        c_idx += 1
                else:
                    val = r.get(h)
                    row_cells[c_idx].text = str(val) if val is not None else ""
                    for run in row_cells[c_idx].paragraphs[0].runs:
                        run.font.size = Pt(7)
                    c_idx += 1
                    
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()
