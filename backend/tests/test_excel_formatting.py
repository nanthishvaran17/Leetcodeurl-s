import io
import openpyxl
import pytest
from backend.exporters.excel_exporter import export_excel_from_dataset
from backend.routes.hr_candidate_finder import generate_hr_candidate_finder_excel

@pytest.fixture
def mock_dataset():
    return {
        "contestName": "Weekly Contest 517",
        "sessionDate": "30.08.2026",
        "snapshotId": "SNAPSHOT_517_TEST",
        "rows": [
            {
                "reg_no": "732224CC031",
                "name": "NANTHISH S",
                "dept": "CSE(CS)",
                "year": "III Year",
                "batch": "2024-2028",
                "section": "A",
                "username": "nanthishvaran_07",
                "status": "PUBLIC_ATTENDED",
                "is_virtual": False,
                "q1": 1,
                "q2": 1,
                "q3": 0,
                "q4": 0,
                "score": 7,
                "rating": 1708.49,
                "rank": 5089,
                "performance_score": 78,
                "acceptance_rate": 62.5,
                "readiness": "On Track",
                "risk": "Safe",
                "trend": "UP"
            },
            {
                "reg_no": "732224CI020",
                "name": "KIRUTHIKA K",
                "dept": "CSE(IOT)",
                "year": "III Year",
                "batch": "2024-2028",
                "section": "B",
                "username": "kiruthika_k",
                "status": "NOT_ATTENDED",
                "is_virtual": False,
                "q1": 0,
                "q2": 0,
                "q3": 0,
                "q4": 0,
                "score": 0,
                "rating": None,
                "rank": None,
                "performance_score": 20,
                "acceptance_rate": 30.0,
                "readiness": "Developing",
                "risk": "High Risk",
                "trend": "STABLE"
            }
        ]
    }

def test_export_excel_from_dataset_formatting(mock_dataset):
    excel_bytes = export_excel_from_dataset(mock_dataset)
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet_names = wb.sheetnames
    
    # Check key sheets
    assert "Executive Summary" in sheet_names
    assert "Complete Student Roster" in sheet_names
    assert "HR Candidate Finder" in sheet_names
    assert "Contest Performance Matrix" in sheet_names
    assert "Difficulty Analysis" in sheet_names
    assert "Data Quality & Audit" in sheet_names
    assert "Methodology" in sheet_names

    # Verify Times New Roman font across all sheets
    non_tnr_cells = []
    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None and cell.font and cell.font.name:
                    # Allow Consolas for SHA-256 hashes specifically
                    if cell.font.name not in ("Times New Roman", "Consolas"):
                        non_tnr_cells.append((sheet_name, cell.coordinate, cell.font.name))
    
    assert len(non_tnr_cells) == 0, f"Found cells not using Times New Roman: {non_tnr_cells[:5]}"

def test_generate_hr_candidate_finder_excel_formatting(mock_dataset):
    candidates = mock_dataset["rows"]
    excel_bytes = generate_hr_candidate_finder_excel(candidates=candidates, filters_desc="Test Cohort")
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet_names = wb.sheetnames
    assert "HR Candidate Finder" in sheet_names

    non_tnr_cells = []
    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None and cell.font and cell.font.name:
                    if cell.font.name not in ("Times New Roman", "Consolas"):
                        non_tnr_cells.append((sheet_name, cell.coordinate, cell.font.name))
    
    assert len(non_tnr_cells) == 0, f"Found cells not using Times New Roman in HR exporter: {non_tnr_cells[:5]}"
